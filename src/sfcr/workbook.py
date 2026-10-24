"""Read a workbook as labelled rows: label, unit, visibility, formula and stored value.

Nothing here knows which figure is which; that is the catalogue's job (locate.py).
"""

import ast
import operator
import re
import unicodedata
from dataclasses import dataclass, field
from decimal import Decimal
from pathlib import Path

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter

from .money import plain, to_decimal

HEADER_SCAN_ROWS = 30
REF = r"\$?([A-Z]{1,3})\$?([0-9]+)"
OPS = {ast.Add: operator.add, ast.Sub: operator.sub, ast.Mult: operator.mul, ast.Div: operator.truediv}


def normalize(text) -> str:
    """Lowercase, no diacritics (ș/ş, ț/ţ alike), punctuation as single spaces."""
    t = unicodedata.normalize("NFKD", str(text))
    t = "".join(c for c in t if not unicodedata.combining(c)).lower()
    return " ".join(re.sub(r"[^0-9a-z]+", " ", t).split())


def label_keys(text) -> set[str]:
    """The label normalised, with and without a trailing parenthesis."""
    keys = {normalize(text)}
    stripped = re.sub(r"\s*\([^()]*\)\s*$", "", str(text))
    if stripped.strip():
        keys.add(normalize(stripped))
    return keys - {""}


class Unsupported(Exception):
    pass


@dataclass
class Value:
    number: Decimal | None  # as stored in the cell, before the sheet's unit
    how: str  # how it was read, for the reason
    problem: str | None = None  # why there is no usable number
    conflict: str | None = None  # stored value differs from the recalculated formula
    number_format: str = ""
    empty: bool = False


@dataclass
class Row:
    row: int
    label: str
    keys: set[str]
    notes: list[str]


@dataclass
class Sheet:
    name: str
    title: str
    ws: object
    wv: object
    header_row: int | None = None
    columns: dict = field(default_factory=dict)  # "current"/"prior" -> (column index, header text)
    unit_text: str | None = None
    factor: Decimal | None = None
    unit_problem: str | None = None
    rows: list = field(default_factory=list)

    def coord(self, row: int, col: int) -> str:
        return f"{get_column_letter(col)}{row}"

    def value(self, row: int, col: int, _depth: int = 0) -> Value:
        coord = self.coord(row, col)
        cell = self.ws.cell(row, col)
        raw, fmt_ = cell.value, cell.number_format or ""
        if raw is None:
            return Value(None, f"{coord} goală", problem=f"celula {self.name}!{coord} este goală", number_format=fmt_,
                         empty=True)
        if isinstance(raw, str) and raw.startswith("="):
            return self._formula(row, col, raw, _depth, fmt_)
        if isinstance(raw, (int, float)) and not isinstance(raw, bool):
            return Value(to_decimal(raw), f"valoare introdusă {plain(to_decimal(raw))}", number_format=fmt_)
        if isinstance(raw, str):
            return Value(None, "text", problem=f"celula {self.name}!{coord} conține text («{raw}»), nu un număr",
                         number_format=fmt_)
        return Value(None, "tip", problem=f"celula {self.name}!{coord} are un tip neașteptat ({type(raw).__name__})",
                     number_format=fmt_)

    def value_at(self, coord: str) -> Value:
        m = re.fullmatch(REF, coord.upper())
        return self.value(int(m.group(2)), column_index_from_string(m.group(1)))

    def _formula(self, row, col, formula, depth, fmt_) -> Value:
        coord = self.coord(row, col)
        cached = self.wv.cell(row, col).value
        try:
            computed, why = self._evaluate(formula, depth), None
        except Unsupported as e:
            computed, why = None, str(e)
        if isinstance(cached, (int, float)) and not isinstance(cached, bool):
            stored = to_decimal(cached)
            if computed is None:
                return Value(stored, f"formulă {formula}, valoare salvată {plain(stored)} (formula nu a putut fi "
                                     f"recalculată: {why})", number_format=fmt_)
            if abs(stored - computed) > Decimal("1e-6") * max(Decimal(1), abs(computed)):
                return Value(stored, f"formulă {formula}", number_format=fmt_,
                             conflict=f"{self.name}!{coord}: valoarea salvată {plain(stored)} diferă de formula "
                                      f"{formula} recalculată = {plain(computed)}")
            return Value(stored, f"formulă {formula} = {plain(computed)}, egală cu valoarea salvată", number_format=fmt_)
        if cached is None and computed is not None:
            return Value(computed, f"formulă {formula} fără valoare salvată, recalculată = {plain(computed)}",
                         number_format=fmt_)
        if cached is None:
            return Value(None, "formulă", number_format=fmt_,
                         problem=f"{self.name}!{coord}: formula {formula} nu are valoare salvată și nu poate fi "
                                 f"recalculată ({why})")
        return Value(None, "formulă", number_format=fmt_,
                     problem=f"{self.name}!{coord}: formula {formula} are valoarea salvată «{cached}»")

    def _evaluate(self, formula: str, depth: int) -> Decimal:
        """Only + - * / on cells of this sheet and SUM(range). Anything else is not evaluated."""
        if depth > 20:
            raise Unsupported("referință circulară sau prea adâncă")
        expr = formula[1:].strip()
        if "!" in expr:
            raise Unsupported("trimite la altă foaie")

        def expand(m):
            c1, r1, c2, r2 = m.group(1), int(m.group(2)), m.group(3), int(m.group(4))
            cells = [f"{get_column_letter(c)}{r}"
                     for r in range(r1, r2 + 1)
                     for c in range(column_index_from_string(c1), column_index_from_string(c2) + 1)]
            return "(" + "+".join(cells) + ")"

        expr = re.sub(rf"SUM\(\s*{REF}\s*:\s*{REF}\s*\)", expand, expr, flags=re.I)
        expr = re.sub(REF, lambda m: f"_{m.group(1)}{m.group(2)}", expr)
        try:
            tree = ast.parse(expr, mode="eval")
        except SyntaxError:
            raise Unsupported(f"expresie necunoscută «{formula}»")

        def ev(node):
            if isinstance(node, ast.Expression):
                return ev(node.body)
            if isinstance(node, ast.BinOp) and type(node.op) in OPS:
                a, b = ev(node.left), ev(node.right)
                if isinstance(node.op, ast.Div) and b == 0:
                    raise Unsupported("împărțire la zero")
                return OPS[type(node.op)](a, b)
            if isinstance(node, ast.UnaryOp) and isinstance(node.op, (ast.USub, ast.UAdd)):
                v = ev(node.operand)
                return -v if isinstance(node.op, ast.USub) else v
            if isinstance(node, ast.Constant) and isinstance(node.value, (int, float)):
                return to_decimal(node.value)
            if isinstance(node, ast.Name) and node.id.startswith("_"):
                m = re.fullmatch(r"_([A-Z]+)([0-9]+)", node.id)
                r, c = int(m.group(2)), column_index_from_string(m.group(1))
                if self.ws.cell(r, c).value is None:
                    return Decimal(0)  # as Excel does
                v = self.value(r, c, depth + 1)
                if v.number is None or v.conflict:
                    raise Unsupported(f"{self.coord(r, c)}: {v.problem or v.conflict}")
                return v.number
            raise Unsupported(f"expresie necunoscută «{formula}»")

        return ev(tree)


@dataclass
class Workbook:
    path: Path
    sheets: list  # visible sheets that may be used
    ignored: list  # (sheet name, reason)


def _title(ws) -> str:
    for row in ws.iter_rows(min_row=1, max_row=5):
        for c in row:
            if isinstance(c.value, str) and c.value.strip() and not c.value.startswith("="):
                return c.value.strip()
    return ""


def _header_match(text: str, aliases: list[str]) -> bool:
    n = normalize(text)
    return any(n == a or n.startswith(a + " ") for a in aliases)


def load(path: Path, run_cfg: dict, year: int) -> Workbook:
    wb_f = openpyxl.load_workbook(path, data_only=False)
    wb_v = openpyxl.load_workbook(path, data_only=True)
    subst = {"year": year, "prior": year - 1}
    current = [normalize(a.format(**subst)) for a in run_cfg["columns"]["current"]]
    prior = [normalize(a.format(**subst)) for a in run_cfg["columns"]["prior"]]
    ignore = [normalize(t) for t in run_cfg["sheets"]["ignore_if_title_contains"]]
    units = {normalize(k): Decimal(v) for k, v in run_cfg["units"].items()}
    unit_re = re.compile(r"(?:^|\bin )(" + "|".join(sorted(map(re.escape, units), key=len, reverse=True)) + r")(?: |$)")

    sheets, ignored = [], []
    for ws in wb_f.worksheets:
        title = _title(ws)
        if ws.sheet_state != "visible":
            ignored.append((ws.title, f"foaie ascunsă (titlu «{title}»)"))
            continue
        hit = next((t for t in ignore if t in normalize(title)), None)
        if hit:
            ignored.append((ws.title, f"titlul «{title}» arată că foaia nu se folosește"))
            continue
        sh = Sheet(ws.title, title, ws, wb_v[ws.title])

        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, HEADER_SCAN_ROWS)):
            cur = [c for c in row if isinstance(c.value, str) and _header_match(c.value, current)]
            if cur:
                sh.header_row = cur[0].row
                sh.columns["current"] = (cur[0].column, cur[0].value)
                pri = [c for c in row if isinstance(c.value, str) and _header_match(c.value, prior)]
                if pri:
                    sh.columns["prior"] = (pri[0].column, pri[0].value)
                break

        # Unit: from the header area (rows above and including the column headers).
        found, texts = {}, []
        for row in ws.iter_rows(min_row=1, max_row=sh.header_row or min(ws.max_row, 5)):
            for c in row:
                if isinstance(c.value, str):
                    n = normalize(c.value)
                    m = unit_re.search(n)
                    if m:
                        found[m.group(1)] = c
                    elif n.startswith("valori"):
                        texts.append(c.value)
        if len(found) == 1:
            key, cell = next(iter(found.items()))
            sh.factor, sh.unit_text = units[key], f"{cell.value} ({sh.name}!{cell.coordinate})"
        elif found:
            sh.unit_problem = "antetul foii indică mai multe unități: " + ", ".join(
                f"«{c.value}» ({c.coordinate})" for c in found.values())
        elif texts:
            sh.unit_problem = f"unitatea din antet nu e recunoscută: «{texts[0]}»"
        else:
            sh.unit_problem = "foaia nu indică unitatea valorilor"

        if sh.header_row:
            value_cols = {col for col, _ in sh.columns.values()}
            first_value_col = min(value_cols)
            for r in range(sh.header_row + 1, ws.max_row + 1):
                labels = [ws.cell(r, c).value for c in range(1, first_value_col)]
                labels = [v.strip() for v in labels if isinstance(v, str) and v.strip() and not v.startswith("=")]
                if not labels:
                    continue
                notes = []
                for c in range(first_value_col, ws.max_column + 1):
                    v = ws.cell(r, c).value
                    if c not in value_cols and isinstance(v, str) and v.strip() and not v.startswith("="):
                        notes.append(f"«{v.strip()}» ({ws.title}!{ws.cell(r, c).coordinate})")
                keys = set().union(*(label_keys(lbl) for lbl in labels))
                sh.rows.append(Row(r, " / ".join(labels), keys, notes))
        sheets.append(sh)
    return Workbook(path, sheets, ignored)
